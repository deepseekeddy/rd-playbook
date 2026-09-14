from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.section import WD_ORIENT
from pathlib import Path


OUT = Path('/Users/eddygao/Documents/ChatGPT/Trae/trae-rd-system/Trae GitHub 团队分工与驾驶舱下一步实施手册.docx')
FONT = 'Heiti SC'
BLUE = '17365D'
LIGHT_BLUE = 'DCE6F1'
PALE_BLUE = 'EEF4F9'
GRAY = 'F2F2F2'
BORDER = 'D9D9D9'
WHITE = 'FFFFFF'
BLACK = '000000'


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'start' if edge == 'left' else 'end' if edge == 'right' else edge
        el = borders.find(qn(f'w:{tag}'))
        if el is None:
            el = OxmlElement(f'w:{tag}')
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), BORDER)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_run_font(run, size=None, bold=None, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(p, after=5, before=0, line=1.18):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = line
    for r in p.runs:
        set_run_font(r)


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    format_paragraph(p)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.38 + 0.25 * level)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        r = p.add_run(f'•  {item}')
        set_run_font(r)
        format_paragraph(p, after=3)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        r = p.add_run(item)
        set_run_font(r)
        format_paragraph(p, after=4)


def add_code(doc, text):
    for line in text.strip('\n').split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(line)
        set_run_font(r, size=9, color='333333')


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        set_run_font(r, size=font_size, bold=True, color=WHITE)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    r = p.add_run(text)
    set_run_font(r, bold=True)
    return p


def add_page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles['Normal']
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.18

for name, size, before, after in [('Title', 25, 0, 12), ('Heading 1', 17, 15, 7), ('Heading 2', 13, 10, 5), ('Heading 3', 11, 7, 4)]:
    st = styles[name]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ['List Bullet', 'List Bullet 2', 'List Number']:
    st = styles[list_name]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.font.size = Pt(10.5)

# Footer
footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run('Trae GitHub 团队协同实施手册  ')
set_run_font(r, size=8, color='666666')
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
p._p.append(fld)

# Cover
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run('Trae GitHub 团队分工与驾驶舱下一步实施手册')
set_run_font(r, size=25, bold=True)

p = doc.add_paragraph()
r = p.add_run('适用于已连接 GitHub 的十人研发团队')
set_run_font(r, size=14, color='333333')
format_paragraph(p, after=18)

add_text(doc, '这份手册从你当前“GitHub 已连接”的状态继续执行。目标是让每项工作在 GitHub 中有唯一负责人、明确状态、审核证据和可聚合字段，并让 Trae 成为执行入口和 AI 助手。')
add_text(doc, '实施结论：GitHub Projects 负责团队任务和驾驶舱；GitHub Issues 与 Sub-issues 负责工作分解；Assignee 负责当前行动；CODEOWNERS 与分支保护负责代码审核；Trae 负责读取项目上下文、完成开发、生成 PR 和辅助检查。')

add_heading(doc, '本周完成标准', 2)
add_bullets(doc, [
    '十名成员全部映射到 GitHub 账号、团队角色和负责目录。',
    '创建一个组织级 GitHub Project，并建立字段、视图和自动化。',
    '选一个真实需求拆成父 Issue 和 5 至 8 个 Sub-issues，每个子任务只有一个 Assignee。',
    '主分支启用 PR、Reviewer、状态检查和 CODEOWNERS 门禁。',
    '驾驶舱能回答每个人在做什么、什么被阻塞、什么待审核、迭代是否有风险。'
])

add_heading(doc, '阅读顺序', 2)
add_table(doc, ['阶段', '你要完成的结果', '建议时间'], [
    ['第一步', '核对能力边界与人员账号', '30 分钟'],
    ['第二步', '创建 GitHub Project 和字段', '60 分钟'],
    ['第三步', '建立个人视图和驾驶舱', '45 分钟'],
    ['第四步', '配置 Issue 模板和任务分配', '90 分钟'],
    ['第五步', '配置审核与分支保护', '60 分钟'],
    ['第六步', '用真实需求试跑', '1 个迭代'],
], widths=[1.2, 4.5, 1.2])

add_page_break(doc)

add_heading(doc, '一 当前能力边界和修正', 1)
add_text(doc, '现有手册已经覆盖 GitHub 授权、分支开发、提交和 PR，但把若干能力描述成“自动双向同步”。当前官方资料明确说明 TraeWork Code 模式可以连接 GitHub、创建 PR 和检查 PR。Issue 创建、GitHub Projects 字段更新、自动分派、自动合并和驾驶舱同步，应通过 GitHub 原生能力、GitHub Actions 或经过验证的 MCP/API 实现，不能仅凭 Trae 对话假设已经完成。')

add_table(doc, ['事项', '正确做法', '不要假设'], [
    ['任务主记录', 'GitHub Issue 与 Project 字段', 'Trae 对话就是团队任务系统'],
    ['个人责任', 'Issue Assignee；每项只设一个主责', '按角色或群聊口头分工'],
    ['代码审核', 'Requested Reviewer、CODEOWNERS、规则集', 'Trae 说审核通过即可合并'],
    ['状态更新', 'Project 字段加 GitHub 自动化或 Actions', '所有状态天然双向实时同步'],
    ['最终合并', 'GitHub 页面或受控合并队列', 'AI 自动合并所有通过的 PR'],
    ['密钥', 'OAuth、GitHub App 或 Actions Secret', '从 remote URL 提取或打印 Token'],
], widths=[1.25, 3.35, 2.3], font_size=8.8)

add_heading(doc, '推荐的系统分工', 2)
add_table(doc, ['系统', '负责内容'], [
    ['GitHub Projects', '迭代、优先级、状态、负责人、时间、风险和驾驶舱'],
    ['GitHub Issues', '需求、开发、测试、文档、发布等具体工作项'],
    ['GitHub Pull Requests', '代码差异、Reviewer、CI、讨论和合并证据'],
    ['Trae', '理解上下文、开发、自测、创建 PR、辅助评审和生成状态摘要'],
    ['飞书', '客户沟通、正式 PRD/TSD、会议、主动通知；只保存 GitHub 链接和版本'],
], widths=[1.55, 5.4])

add_page_break(doc)

add_heading(doc, '今天先准备的信息', 2)
add_table(doc, ['姓名', 'GitHub 用户名', '角色', '负责目录或模块', '备份人'], [
    ['成员 1', '@username', 'PM', '产品需求', '成员 2'],
    ['成员 2', '@username', '设计', 'design 或 Figma', '成员 1'],
    ['成员 3', '@username', 'TL', '全局架构与核心服务', '成员 4'],
    ['成员 4 至 8', '@username', '开发', '具体目录或服务', '另一名开发'],
    ['成员 9', '@username', 'QA', '测试与验收', 'TL'],
    ['成员 10', '@username', 'SRE', '部署与监控', 'TL'],
], widths=[1.05, 1.25, 0.8, 2.8, 1.0], font_size=8.3)

add_heading(doc, '二 创建团队驾驶舱', 1)
add_heading(doc, '步骤一 创建组织级 Project', 2)
add_steps(doc, [
    '打开 GitHub Organization，进入 Projects，点击 New project。',
    '选择空白 Project 或 Team planning 模板，命名为 R&D Delivery。',
    '在 Project 的 README 中写明状态定义、字段说明、Owner 和协作规则。',
    '将项目可见范围设为团队可见；只给项目管理员编辑工作流和字段的权限。',
    '在 Project 菜单进入 Workflows，打开 Auto-add to project，选择已经连接的仓库。'
])

add_heading(doc, '步骤二 创建字段', 2)
add_table(doc, ['字段', '类型', '建议值或规则', '驾驶舱用途'], [
    ['Status', 'Single select', 'Intake、Ready、In Progress、In Review、Ready for QA、In QA、Ready for Release、Blocked、Done', '流程位置'],
    ['Iteration', 'Iteration', '每两周一个迭代', '当前迭代和燃尽'],
    ['Priority', 'Single select', 'P0、P1、P2、P3', '排序'],
    ['Workstream', 'Single select', 'Product、Design、Frontend、Backend、QA、Ops', '职能负载'],
    ['Type', 'Single select', 'Epic、Feature、Task、Bug、Review、Change', '工作类型'],
    ['Risk', 'Single select', 'Low、Medium、High', '审核深度'],
    ['Estimate', 'Number', '1、2、3、5、8', '容量'],
    ['Start date', 'Date', '计划开始日', 'Roadmap'],
    ['Target date', 'Date', '承诺完成日', '逾期'],
    ['Health', 'Single select', 'On track、At risk、Off track', '管理状态'],
], widths=[1.0, 0.85, 3.2, 1.75], font_size=8.1)

add_text(doc, '人员字段使用 GitHub 原生 Assignees。不要为十个人各建一个自定义状态，也不要用标签代替项目状态。非代码审核建议单独创建 Review 类型 Issue 并指派审核人；代码审核使用 PR 的 Requested reviewers 和 CODEOWNERS。')

add_heading(doc, '步骤三 配置状态定义', 2)
add_table(doc, ['状态', '进入条件', '退出条件'], [
    ['Intake', '新需求或新缺陷已创建', '完成分诊，明确价值和 DRI'],
    ['Ready', '输入、验收标准、Assignee 齐全', 'Owner 开始执行'],
    ['In Progress', '工作正在执行', '已有可审核产物'],
    ['In Review', 'PR 或文档已提交', 'Reviewer 批准或退回'],
    ['Ready for QA', '代码合并且候选构建成功', 'QA 接收版本'],
    ['In QA', '环境、用例和版本明确', '测试通过或退回'],
    ['Ready for Release', '测试通过、发布证据齐全', 'Go/No-Go 决策'],
    ['Blocked', '存在明确阻塞', '阻塞解除且记录处理结果'],
    ['Done', '验收完成且证据链接齐全', '无需继续流转'],
], widths=[1.25, 2.85, 2.85], font_size=8.6)

add_page_break(doc)

add_heading(doc, '三 精确到每个人的任务设计', 1)
add_heading(doc, '一项工作只设一个主责', 2)
add_bullets(doc, [
    '一个 Issue 只设置一个主要 Assignee；协作者通过评论、链接的子 Issue 或 Reviewer 表达。',
    'Assignee 对下一步行动负责，父需求的 DRI 对整体交付负责。两者可以不同。',
    '任务进入 In Review 后，原 Assignee 不删除；Reviewer 通过 PR 或独立 Review Issue 体现。',
    '成员休假或 WIP 超限时，TL 明确改派 Assignee，不能只在群里说“谁有空帮一下”。'
])

add_heading(doc, '父 Issue 和 Sub-issues 示例', 2)
add_text(doc, '父 Issue：REQ 2026 042 批量导出订单。父 Issue 由 PM 王敏负责，用于保存目标、范围、成功指标、PRD/TSD/Figma 链接和整体进度。')
add_table(doc, ['Sub-issue', 'Assignee', '输出', '完成条件'], [
    ['PM 需求澄清与 PRD', '王敏', 'PRD v1.0', 'PM、设计、TL、QA 审核完成'],
    ['UX 导出交互与状态', '李婷', 'Figma v1.0', '所有状态与 FR-ID 对齐'],
    ['BE 异步导出任务 API', '张强', 'PR 和接口文档', 'CI 通过、Reviewer 批准'],
    ['FE 导出配置与进度 UI', '赵明', 'PR 和交互自测', 'CI 通过、设计验收'],
    ['QA 用例与回归', '赵雪', '测试报告', 'P0/P1 清零'],
    ['OPS 发布与监控', '周宁', '发布单和看板', '发布验证与观察窗完成'],
], widths=[2.25, 0.85, 1.6, 2.25], font_size=8.4)

add_heading(doc, 'Issue 必填结构', 2)
add_code(doc, '''标题：[BE][REQ-2026-042] 实现异步导出任务 API

Owner：@zhangqiang
输入版本：PRD v1.0 / TSD v1.0
目标：提供创建、查询和下载导出任务的 API
范围外：前端 UI、历史导出迁移
依赖：#101 数据权限规则批准
验收：单元测试、集成测试、10 万行性能验证
输出：PR、API 文档、日志和指标
风险：PII、文件生命周期、重试幂等
截止时间：2026-09-18
关闭方式：PR 使用 Closes #123''')

add_heading(doc, '用 Trae 拆任务的指令', 2)
add_code(doc, '''读取已批准的 PRD、TSD 和当前 GitHub Project。
把 REQ-2026-042 拆成可独立验收的 Sub-issues。
每个 Sub-issue 必须只有一个 Owner、明确输入版本、输出、依赖、Estimate 和完成条件。
按团队模块责任表推荐 Assignee，但先生成草稿，不创建 Issue。
检查每人的当前 In Progress 数量；超过 2 项时给出改派建议。''')

add_heading(doc, '四 建立每个人的工作视图', 1)
add_text(doc, 'GitHub Projects 使用同一批 Issue 创建多个视图。字段只有一份，视图只是不同的过滤和展示方式，因此不会形成多套状态。')

add_table(doc, ['视图', '布局', '过滤或分组', '用途'], [
    ['团队驾驶舱', 'Table', 'Iteration 为当前；按 Status 分组', '管理整体健康度和阻塞'],
    ['我的工作', 'Board', 'assignee:@me 且非 Done；列为 Status', '成员每日使用'],
    ['待我审核', 'Table', 'Review 类型且 assignee:@me，或 GitHub PR review-requested:@me', '审核队列'],
    ['迭代看板', 'Board', '当前 Iteration；列为 Status', '站会和流转'],
    ['负载视图', 'Table', '当前 Iteration；按 Assignee 分组', '发现 WIP 和分配冲突'],
    ['风险与阻塞', 'Table', 'Status=Blocked 或 Health 非 On track', '管理升级'],
    ['发布路线图', 'Roadmap', '显示 Start date 和 Target date', '版本与依赖'],
    ['客户变更', 'Table', 'Type=Change', '追踪范围变更和影响'],
], widths=[1.2, 0.75, 3.0, 2.0], font_size=8.4)

add_heading(doc, '驾驶舱顶部要回答的问题', 2)
add_table(doc, ['问题', '字段或图表', '红线'], [
    ['每个人在做什么', 'Assignee 分组，显示 Status、Estimate、Target date', '一个人 In Progress 超过 2 项'],
    ['迭代是否能完成', 'Status 按 Estimate 汇总；Iteration 过滤', '到期前两天仍有 High Risk 未评审'],
    ['哪里被阻塞', 'Blocked 视图和阻塞时长', '超过 24 小时无 Unblock Owner'],
    ['什么待审核', 'Review Issue 与 Requested reviewers', '超过 8 工作小时'],
    ['质量是否达标', 'CI、Bug、Ready for QA、In QA', 'P0/P1 未关闭'],
    ['范围是否变化', 'Change Issue、版本基线和关联父 Issue', 'Approved 后变更未重开门禁'],
], widths=[1.55, 3.15, 2.25], font_size=8.2)

add_heading(doc, '建议图表', 2)
add_bullets(doc, [
    '按 Status 统计当前迭代的 Issue 数量或 Estimate。',
    '按 Assignee 统计未完成 Estimate，识别个人负载差异。',
    '按 Workstream 统计 Ready、In Progress 和 Blocked，识别职能瓶颈。',
    '按 Risk 和 Health 统计高风险、偏离计划的事项。',
    'Bug 按严重度和状态统计；发布前单独查看 P0/P1。'
])

add_text(doc, '驾驶舱用于发现流程问题，不用于按代码行数、提交数或 AI 生成量评价个人。')

add_page_break(doc)

add_heading(doc, '五 审核和合并机制', 1)
add_heading(doc, '配置 GitHub Teams', 2)
add_table(doc, ['Team', '建议成员', '用途'], [
    ['product-design', 'PM、设计', 'PRD、设计与体验相关审核'],
    ['frontend', '前端成员', '前端目录 CODEOWNERS'],
    ['backend', '后端成员', '服务端目录 CODEOWNERS'],
    ['qa', 'QA、自动化负责人', '测试策略与质量审核'],
    ['maintainers', 'TL、资深开发、SRE', '仓库规则和高风险审核'],
], widths=[1.5, 2.25, 3.2])

add_heading(doc, 'CODEOWNERS 示例', 2)
add_code(doc, '''# 全局默认审核
*                         @your-org/maintainers

# 目录责任
/apps/web/                @your-org/frontend
/tests/                   @your-org/qa
/.github/                 @your-org/maintainers
/deploy/                  @your-org/maintainers
/docs/                    @your-org/product-design @your-org/maintainers''')

add_heading(doc, '主分支规则集', 2)
add_steps(doc, [
    '进入 Repository Settings，打开 Rules 或 Branch protection，目标分支选择 main。',
    '启用 Require a pull request before merging，禁止直接推送 main。',
    '普通变更至少 1 个批准；High Risk 或核心目录建议 2 个批准。',
    '启用 Require review from Code Owners。',
    '启用 Dismiss stale approvals 或要求最新可审核提交由非提交者批准。',
    '启用 Require status checks，选择 lint、unit、integration、security 等实际存在的检查。',
    '启用 Require conversation resolution；必要时启用 merge queue。',
    '禁止删除和 force push；只给极少数管理员紧急绕过权限，并记录原因。'
])

add_heading(doc, '审核状态如何进入驾驶舱', 2)
add_bullets(doc, [
    '开发 Issue 在开始时为 In Progress。',
    'PR 标记 Ready for review 后，Issue 或 PR 项进入 In Review。',
    'CODEOWNERS 自动请求对应团队审核；Trae 可以辅助检查，但不能写入真人批准。',
    '审批和 CI 全部通过后才允许合并；PR 合并后，链接的 Issue 通过 Closes 关闭。',
    '需要 QA 的需求不要在代码合并后直接将父需求设为 Done，应进入 Ready for QA。'
])

add_heading(doc, '六 自动化配置', 1)
add_heading(doc, '先启用 GitHub Projects 内置工作流', 2)
add_table(doc, ['触发', '建议动作'], [
    ['Issue 或 PR 符合仓库和标签过滤', '自动添加到 R&D Delivery Project'],
    ['Item 新增到 Project', 'Status 设置为 Intake 或 Ready'],
    ['Issue 重新打开', 'Status 恢复到 Ready 或 In Progress'],
    ['Issue 或 PR 关闭或合并', '根据工作类型更新为 Done；涉及 QA 时由 Action 改为 Ready for QA'],
    ['Item 达到 Done 且超过保留期', '自动归档'],
], widths=[2.55, 4.4])

add_text(doc, 'GitHub 内置 Auto-add 只会处理创建后或更新后匹配过滤器的事项，不会自动补入所有历史事项。配置完成后，历史 Issue 需要手工批量加入。')

add_heading(doc, '再用 GitHub Actions 处理团队规则', 2)
add_table(doc, ['事件', '自动动作', '人工门禁'], [
    ['新 Issue', '校验模板字段、加入 Project、设置默认 Status', 'PM/TL 分诊'],
    ['Issue 被指派', 'Status 从 Intake 进入 Ready', 'Assignee 接受任务'],
    ['PR Ready for review', '加入 Project、Status=In Review、记录日期', 'Reviewer 审核'],
    ['CI 失败', 'Health=At risk，通知 Author', '修复或批准豁免'],
    ['PR 合并', '关闭开发 Issue，父需求进入 Ready for QA', 'QA 接受版本'],
    ['Target date 超期', 'Health=Off track，通知 Assignee 和 DRI', '改派或调整计划'],
    ['客户 Change 创建', '标记受影响 Issue，生成影响分析任务', 'PM/TL 批准变更'],
], widths=[1.65, 3.25, 2.05], font_size=8.3)

add_heading(doc, '主动通知', 2)
add_text(doc, 'GitHub 原生通知负责 Assignee、Mention、Requested reviewer 和 CODEOWNERS。若团队希望飞书主动提醒，使用 GitHub Actions 或轻量工作流服务监听 Issue、PR、Review、CI 和 Project 事件，再通过飞书机器人发送消息。Trae 本身不应作为后台事件监听器。')
add_table(doc, ['级别', '渠道', '示例'], [
    ['Action', 'GitHub 通知加飞书私聊', '任务指派、阻塞处理、CI 失败'],
    ['Review', 'GitHub Review Request 加飞书私聊', 'PRD、TSD、PR 审核'],
    ['FYI', '飞书项目日报', '排期调整、已批准变更'],
    ['Escalation', '飞书私聊 Assignee 与 TL', '审核超时、阻塞超过 24 小时'],
], widths=[1.05, 2.5, 3.4])

add_heading(doc, '通知示例', 2)
add_code(doc, '''[需审核｜8 工作小时内]
REQ-2026-042 订单导出 API PR #231
你被通知的原因：CODEOWNER /services/order-export
需要行动：检查权限、幂等、审计日志和测试证据
当前状态：In Review
CI：unit 通过，integration 通过，performance 未执行
PR：https://github.com/ORG/REPO/pull/231
逾期升级：TL''')

add_page_break(doc)

add_heading(doc, '七 团队在 Trae 中的每日操作', 1)
add_heading(doc, '成员开始工作', 2)
add_steps(doc, [
    '打开 GitHub Project 的我的工作视图，只选择一个 Ready 事项进入 In Progress。',
    '在 TraeWork Code 模式选择目标仓库和分支，任务名称包含 Issue 编号。',
    '在第一条指令中引用 Issue、PRD/TSD 版本和验收标准。',
    '让 Trae 先输出实施计划；成员确认后再修改代码。',
    '完成后运行项目规则要求的检查，推送分支并创建 PR。',
    '在 GitHub 确认 Requested reviewers、CI 和 Project 状态。'
])

add_heading(doc, '开发提示词', 2)
add_code(doc, '''实现 GitHub Issue #123。
事实源：Issue #123、PRD v1.0、TSD v1.0、AGENTS.md。
先检查输入是否齐全并给出计划；确认后再修改。
只处理 Issue 范围，补充单元和集成测试。
完成后列出修改文件、实际运行的测试、未执行检查、风险和 PR 描述。
不要合并 PR，不要更新父需求为 Done。''')

add_heading(doc, 'Reviewer 开始审核', 2)
add_steps(doc, [
    '打开待我审核视图或 GitHub Review Request。',
    '核对 Issue、PRD/TSD 版本和变更范围。',
    '点击 TraeWork 的 AI 检查 PR，或调用 code-review Skill 获取辅助意见。',
    'Reviewer 自己检查关键逻辑和测试证据，并在 GitHub 选择 Approve 或 Request changes。',
    '有新提交时重新确认最新差异；满足规则后再合并。'
])

add_heading(doc, '站会只看三个视图', 2)
add_bullets(doc, [
    '风险与阻塞：先处理 Blocked 和 Off track。',
    '待审核：处理超过 SLA 的 Review。',
    '当前迭代：确认每人 In Progress 不超过 2 项，以及本周目标是否变化。'
])

add_heading(doc, '项目负责人每日摘要提示词', 2)
add_code(doc, '''基于 GitHub Project 当前迭代生成团队摘要。
按成员列出 In Progress、Blocked、待审核和今天到期事项。
指出 WIP 超过 2、阻塞超过 24 小时、审核超过 8 工作小时、High Risk 未评审的事项。
每条必须包含 Issue 或 PR 链接、当前责任人和下一动作。
不要用提交数或代码行数评价成员。''')

add_page_break(doc)

add_heading(doc, '八 用一个真实需求完成试跑', 1)
add_heading(doc, '第一天 创建和分解', 2)
add_steps(doc, [
    '选择一个中等规模、非紧急的真实需求，创建父 Issue。',
    '使用 Trae 生成 5 至 8 个 Sub-issue 草稿。',
    'PM 和 TL 确认范围、依赖、Estimate、Assignee 和 Reviewer。',
    '把所有事项加入当前 Iteration，并设置 Priority、Risk、Target date。',
    '检查团队负载视图，解决 WIP 超限和单点负责人。'
])

add_heading(doc, '第二至四天 执行和审核', 2)
add_steps(doc, [
    '各 Owner 从我的工作视图领取一项 Ready 任务。',
    '开发在 Trae 中按 Issue 开分支、实现、自测并创建 PR。',
    'CODEOWNERS 自动请求 Reviewer；Actions 更新 In Review。',
    'Reviewer 在 GitHub 做正式决定；CI 失败不得进入下一阶段。',
    '每天站会只处理阻塞、待审核和风险，不逐人朗读任务。'
])

add_heading(doc, '第五天 测试和复盘', 2)
add_steps(doc, [
    '代码合并后父需求进入 Ready for QA，QA 明确接收候选版本。',
    '测试结果、Bug 和证据关联到父 Issue；P0/P1 未关闭不得发布。',
    '完成发布后再关闭父 Issue，不把 PR 合并等同于需求完成。',
    '复盘驾驶舱字段是否真实反映工作；删除无人在意的字段和通知。'
])

add_heading(doc, '试跑验收', 2)
add_table(doc, ['检查项', '通过标准'], [
    ['人员责任', '所有开放 Issue 有一个明确 Assignee；Reviewer 可追踪'],
    ['工作分解', '父 Issue 有 Sub-issues，依赖和完成条件清楚'],
    ['状态真实性', 'In Review 有 PR 或文档；Done 有验收证据'],
    ['审核约束', '主分支不能直接推送；CODEOWNERS 和 CI 生效'],
    ['驾驶舱', '能按 Assignee、Status、Risk、Iteration 查看'],
    ['通知', '行动类即时通知；FYI 汇总；超时可升级'],
], widths=[1.45, 5.5])

add_page_break(doc)

add_heading(doc, '九 七天实施清单', 1)
add_table(doc, ['日期', 'Owner', '动作', '产出'], [
    ['第 1 天', 'GitHub 管理员', '邀请成员、建立 Teams、核对权限', '账号和责任映射'],
    ['第 2 天', 'PM 与 TL', '创建组织 Project、字段和状态', 'R&D Delivery'],
    ['第 3 天', 'PM 与 TL', '创建八个视图和基础图表', '团队驾驶舱'],
    ['第 4 天', 'TL', '提交 Issue 模板、PR 模板、CODEOWNERS', '标准化入口'],
    ['第 5 天', 'GitHub 管理员', '配置 ruleset、CI 和内置 Project workflow', '强制审核与自动状态'],
    ['第 6 天', '全员', '用一个真实需求试跑', '父 Issue、Sub-issues、PR'],
    ['第 7 天', 'PM、TL、QA', '复盘字段、通知、SLA 和权限', 'v1.0 协作规则'],
], widths=[0.9, 1.3, 3.35, 1.4], font_size=8.5)

add_heading(doc, '第一步现在就做', 2)
add_text(doc, '今天先不要配置复杂自动化。请完成下面三项，然后再进入 Project 字段配置。')
add_steps(doc, [
    '整理十名成员的 GitHub 用户名、角色、负责目录、备份人。',
    '确认仓库属于个人账号还是 GitHub Organization；如果是个人仓库，优先迁移到 Organization 后再建立团队 Project。',
    '在 GitHub 新建组织级 R&D Delivery Project，并把一个真实 Issue 加进去。'
])

add_heading(doc, '完成后记录这些信息', 2)
add_code(doc, '''GitHub Organization：
Repository：
Project URL：
默认分支：
当前 CI 检查名称：
十名成员 GitHub 用户名：
当前是否有 CODEOWNERS：是 / 否
当前是否有 ruleset 或 branch protection：是 / 否
计划使用两周还是一周 Iteration：''')

add_heading(doc, '下一轮配置所需输入', 2)
add_text(doc, '拿到上述信息后，可以继续生成与你们真实组织、仓库、成员和检查名称完全匹配的 Issue Form、PR 模板、CODEOWNERS、GitHub Actions 和驾驶舱字段配置。')

add_heading(doc, '十 官方参考', 1)
add_bullets(doc, [
    'TRAE GitHub 集成：https://docs.trae.cn/work_github-integration',
    'GitHub Projects：https://docs.github.com/en/issues/planning-and-tracking-with-projects/learning-about-projects/about-projects',
    'GitHub Projects 快速开始：https://docs.github.com/en/issues/planning-and-tracking-with-projects/learning-about-projects/quickstart-for-projects',
    'GitHub Sub-issues：https://docs.github.com/en/issues/tracking-your-work-with-issues/using-issues/adding-sub-issues',
    'GitHub Project 自动添加：https://docs.github.com/en/issues/planning-and-tracking-with-projects/automating-your-project/adding-items-automatically',
    'GitHub CODEOWNERS：https://docs.github.com/en/repositories/managing-your-repositorys-settings-and-features/customizing-your-repository/about-code-owners',
    'GitHub 分支保护：https://docs.github.com/en/repositories/configuring-branches-and-merges-in-your-repository/managing-protected-branches/about-protected-branches'
])

# Normalize all paragraph fonts and keep tables with headings where practical
for p in doc.paragraphs:
    for r in p.runs:
        if not r.font.name:
            set_run_font(r)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
